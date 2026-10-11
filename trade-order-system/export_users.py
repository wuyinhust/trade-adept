import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "用户权限总表"

# ── 样式定义 ──
header_font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
cell_font = Font(name="微软雅黑", size=10)
cell_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin")
)
team_a_fill = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
team_b_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
admin_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

# ── 数据 ──
users = [
    # (序号, 用户名, 密码, 姓名, 角色, 团队, 可见公司, 可访问页面, 能否删除)
    (1, "David", "Fuckyou361021", "David", "管理员 (admin)", "管理层",
     "全部4家公司\n(国润/为语/Grandware HK/Likeing HK)",
     "总览/订单/客户/工厂/出运/利润/单证/系统设置", "是"),
    (2, "Lily", "lily0621", "Lily", "财务 (finance)", "管理层",
     "全部4家公司\n(国润/为语/Grandware HK/Likeing HK)",
     "总览/订单/客户/工厂/出运/利润/单证/系统设置", "否"),
    (3, "Vivian", "vivian8499", "Vivian", "经理 (manager)", "Team A",
     "全部4家公司\n(国润/为语/Grandware HK/Likeing HK)",
     "总览/订单/客户/工厂/出运/利润/单证/系统设置", "否"),
    (4, "Michale", "michale8905", "Michale", "业务员 (salesperson)", "Team A",
     "全部4家公司\n(国润/为语/Grandware HK/Likeing HK)",
     "总览/订单/客户/工厂/出运/利润/单证", "否"),
    (5, "Alex", "alex6686", "Alex", "业务员 (salesperson)", "Team A",
     "2家公司\n(为语/Grandware HK)",
     "总览/订单/客户/工厂/出运/利润/单证", "否"),
    (6, "Anna", "anna9586", "Anna", "单证员 (document)", "Team A",
     "全部4家公司\n(国润/为语/Grandware HK/Likeing HK)",
     "总览/订单/工厂/出运/单证", "否"),
    (7, "Sue", "sue2564", "Sue", "业务员 (salesperson)", "Team A",
     "1家公司\n(Likeing HK)",
     "总览/订单/客户/工厂/出运/利润/单证", "否"),
    (8, "Galen", "galen5907", "Galen", "经理+业务 (manager)", "Team B",
     "2家公司\n(为语/Grandware HK)",
     "总览/订单/客户/工厂/出运/利润/单证/系统设置", "否"),
    (9, "Nicole", "nicole8426", "Nicole", "业务员 (salesperson)", "Team B",
     "2家公司\n(为语/Grandware HK)",
     "总览/订单/客户/工厂/出运/利润/单证", "否"),
    (10, "Jack", "jack9557", "Jack", "单证员 (document)", "Team B",
     "2家公司\n(为语/Grandware HK)",
     "总览/订单/工厂/出运/单证", "否"),
    (11, "Melody", "melody3323", "Melody", "单证员 (document)", "Team B",
     "2家公司\n(为语/Grandware HK)",
     "总览/订单/工厂/出运/单证", "否"),
]

# ── 标题行 ──
title_row = ["序号", "用户名", "密码", "姓名", "角色", "所属团队", "可见公司", "可访问页面", "能否删除订单"]

# 写标题
for col_idx, title in enumerate(title_row, 1):
    cell = ws.cell(row=1, column=col_idx, value=title)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align
    cell.border = thin_border

# 写数据
for row_idx, user in enumerate(users, 2):
    team = user[5]
    if team == "Team A":
        fill = team_a_fill
    elif team == "Team B":
        fill = team_b_fill
    else:
        fill = admin_fill

    for col_idx, value in enumerate(user, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.font = cell_font
        cell.fill = fill
        cell.border = thin_border
        # 可见公司和页面列左对齐
        if col_idx in (7, 8):
            cell.alignment = left_align
        else:
            cell.alignment = cell_align

# ── 列宽 ──
col_widths = [6, 12, 18, 10, 22, 10, 28, 38, 14]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# 行高
ws.row_dimensions[1].height = 28
for r in range(2, len(users) + 2):
    ws.row_dimensions[r].height = 48

# 冻结首行
ws.freeze_panes = "A2"

# ── Sheet 2: 说明 ──
ws2 = wb.create_sheet("权限说明")
notes = [
    ["说明项", "内容"],
    ["管理员 (admin)", "拥有系统全部权限，可查看所有公司和订单，可删除订单"],
    ["财务 (finance)", "可查看所有公司的全部页面，不可删除订单"],
    ["经理 (manager)", "可查看团队的订单、利润分析和系统设置；Team B 经理 Galen 同时承担业务角色"],
    ["业务员 (salesperson)", "可管理自己的订单、客户、工厂出运、利润和单证；不可看系统设置"],
    ["单证员 (document)", "可管理团队订单、工厂和物流单证；不可看利润、客户和系统设置"],
    ["", ""],
    ["公司简称对照", ""],
    ["国润", "Grandware (国润)"],
    ["为语", "为语"],
    ["Grandware HK", "Grandware HK"],
    ["Likeing HK", "Likeing HK"],
    ["", ""],
    ["密码规则", "David 和 Lily 保留原密码；其余人员密码格式：名字全小写 + 4位随机数字"],
    ["导出时间", "2026年7月29日"],
]

for row_idx, row_data in enumerate(notes, 1):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws2.cell(row=row_idx, column=col_idx, value=value)
        if row_idx == 1:
            cell.font = Font(name="微软雅黑", bold=True, size=11, color="FFFFFF")
            cell.fill = header_fill
        else:
            cell.font = Font(name="微软雅黑", size=10)
        cell.alignment = left_align if col_idx == 2 else cell_align
        cell.border = thin_border

ws2.column_dimensions["A"].width = 20
ws2.column_dimensions["B"].width = 55

# ── 保存 ──
output_path = r"D:\Users\workbuddy\2026-07-28-11-24-08\trade-order-system\用户权限总表.xlsx"
wb.save(output_path)
print(f"Excel saved to: {output_path}")
